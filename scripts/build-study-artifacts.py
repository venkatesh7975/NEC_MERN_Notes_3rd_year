"""Export original study content into searchable PDF and an editable study plan.
Run after build-handbook.py. Requires requirements-artifacts.txt.
"""
from pathlib import Path
import json,textwrap,re,sys
from html import escape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Preformatted,PageBreak,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'interview-handbook/data/study-data.json').read_text())
resources=json.loads((ROOT/'interview-handbook/resources/catalog.json').read_text())
out=ROOT/'interview-handbook/downloads';out.mkdir(parents=True,exist_ok=True)
styles=getSampleStyleSheet()
styles['Title'].fontName='Helvetica-Bold';styles['Title'].fontSize=28;styles['Title'].leading=34
styles['Heading1'].textColor=colors.HexColor('#17376b');styles['Heading1'].spaceBefore=16;styles['Heading1'].spaceAfter=10
styles['Heading2'].fontSize=12;styles['Heading2'].leading=16;styles['Heading2'].spaceBefore=12
styles['BodyText'].fontSize=10;styles['BodyText'].leading=15;styles['BodyText'].spaceAfter=8
styles.add(ParagraphStyle(name='CodeSafe',fontName='Courier',fontSize=7.4,leading=10,backColor=colors.HexColor('#eef2f7'),borderPadding=8,spaceBefore=5,spaceAfter=12))
story=[]
def p(text,style='BodyText'): return Paragraph(escape(text),styles[style])
def add(text,style='BodyText'):story.append(p(text,style))
def footer(canvas,doc):
    canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#52627a'))
    canvas.drawString(46,28,'MERN interview handbook | Original practice | 2026-10-08')
    canvas.drawRightString(A4[0]-46,28,str(doc.page));canvas.restoreState()
add('MERN interview handbook','Title')
add('A study and practice reference for product companies and startups')
add('This handbook contains ten revision guides, 100 answered questions classified by learning difficulty, 24 debugging scenarios, and 16 machine coding approaches. Use the repository for runnable projects, mind maps, company playbooks, and the progress tracker.')
add('How to study','Heading1')
add('Read one idea, close the guide, explain it aloud, implement an example, and demonstrate a boundary case. Score recall from 0 to 3: cannot explain, definition only, correct with example, or correct with tradeoff and failure case. Revisit weak answers after 1, 3, 7, and 14 days.')
add('Scope and sources','Heading1')
add('Questions and scenarios are original practice, not leaked company questions. No material guarantees interview selection. External resources retain their licenses; links are included without reproducing external ebooks. The code examples are learning references and may omit a surrounding application setup explicitly provided in the repository.')
story.append(Paragraph('Repository: <link href="https://github.com/venkatesh7975/NEC_MERN_Notes_3rd_year">NEC MERN learning resources</link>',styles['BodyText']))
add('Contents','Heading1')
for title in ['Part 1 Structured revision notes','Part 2 Difficulty based questions and answers','Part 3 Scenario diagnosis','Part 4 Machine coding approaches','Part 5 External learning references']:add(title)
story.append(PageBreak());add('Part 1 Structured revision notes','Title')
for note_index,note in enumerate(data['notes']):
    if note_index: story.append(PageBreak())
    add(note['title'],'Heading1')
    for heading,a,b in note['sections']:add(heading,'Heading2');add(a);add(b)
    add('Worked example','Heading2')
    code='\n'.join('\n'.join(textwrap.wrap(line,width=94,replace_whitespace=False,drop_whitespace=False,subsequent_indent='    ')) or '' for line in note['code'].splitlines())
    story.append(Preformatted(code,styles['CodeSafe']))
    add('Demonstrate understanding','Heading2');add(note['exercise'])
    story.append(Paragraph(f'Primary reference: <link href="{escape(note["source"])}">{escape(note["source"])}</link>',styles['BodyText']))
story.append(PageBreak());add('Part 2 Difficulty based questions and answers','Title')
last=None
for q in data['questions']:
    if q['topic']!=last:add(q['topic'].replace('-',' ').title(),'Heading1');last=q['topic']
    story.append(KeepTogether([p(f'{q["id"]} {q["difficulty"]} {q["question"]}','Heading2'),p(q['answer']),p('Follow-up: '+q['followup'])]))
story.append(PageBreak());add('Part 3 Scenario diagnosis','Title')
for s in data['scenarios']:
    start=len(story)
    add(s['id']+' '+s['symptom'],'Heading2')
    for label,field in [('Evidence','evidence'),('Likely cause','cause'),('Fix','fix'),('Verification','test'),('Common mistake','pitfall')]:add(label+': '+s[field])
    block=story[start:];story[start:]=[KeepTogether(block)]
story.append(PageBreak());add('Part 4 Machine coding approaches','Title')
add('Attempt a build before reading its approach. Times are suggested limits. Score required behavior 30, boundaries and recovery 20, state and API design 15, accessibility 15, verification 10, and explanation 10.')
for m in data['machine']:
    start=len(story)
    add(f'{m["id"]} {m["title"]}','Heading2');add(f'{m["level"]} | {m["minutes"]} minutes | {m["topic"]}')
    add('Required behavior: '+m['requirements']);add('Boundary checks: '+m['edges']);add('Solution approach: '+m['approach'])
    block=story[start:];story[start:]=[KeepTogether(block)]
story.append(PageBreak());add('Part 5 External learning references','Title')
add('Links were reviewed on 2026-10-08. A successful retrieval does not establish unchanged prices or hiring processes. Consult the repository catalog for partial checks.')
for resource in resources:
    start=len(story)
    add(resource['name'],'Heading2');add(resource['topic']+' | '+resource['level']+' | '+resource['access']);add(resource['use'])
    story.append(Paragraph(f'<link href="{escape(resource["url"])}">{escape(resource["url"])}</link>',styles['BodyText']))
    block=story[start:];story[start:]=[KeepTogether(block)]
SimpleDocTemplate(str(out/'mern-interview-handbook.pdf'),pagesize=A4,rightMargin=46,leftMargin=46,topMargin=42,bottomMargin=44,title='MERN interview handbook',author='NEC MERN repository contributors').build(story,onFirstPage=footer,onLaterPages=footer)
if '--pdf-only' in sys.argv:
    print('Created complete PDF handbook.');sys.exit(0)

# A concise editable planning document complements the complete PDF reference.
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(28)
doc.add_heading('MERN interview study plan',0)
doc.add_paragraph('Product company and startup preparation',style='Subtitle')
doc.add_paragraph('Use this editable plan to turn the repository into a weekly practice schedule. It covers the learning sequence, interview recall, timed builds, project evidence, and mock review. The complete PDF handbook provides the answers and worked explanations. Readiness is established by demonstrated behavior and reasoning; this plan does not promise a hiring outcome.')
doc.add_heading('Your preparation target',1)
for line in ['Target role and level ______________________________','Company and job URL ______________________________','Recruiter confirmed rounds ______________________________','Available hours each week ______________________________','Interview date if known ______________________________']:doc.add_paragraph(line)
doc.add_heading('Start with a diagnostic',1)
doc.add_paragraph('Answer five JavaScript questions, complete one 60-minute UI build, and trace an authenticated API write. Record where you needed help, which behavior failed, and what you could explain. Choose the next week from that evidence.')
doc.add_page_break();doc.add_heading('Twelve week learning sequence',1)
weeks=[('1 and 2','Web HTML CSS and accessibility','Trace a request, create a semantic form, and build a responsive layout. Verify keyboard use, 320px width, and 200 percent zoom.'),('3 and 4','JavaScript async and algorithms','Explain closures and event ordering. Implement debounce, an emitter, an LRU, and a promise pool. Show invalid-input and failure behavior.'),('5 and 6','React and machine coding','Model state ownership, stable keys, effects, and remote states. Complete search and table exercises under a timer.'),('7 and 8','Node Express and MongoDB','Build a protected API. Explain middleware, validation, indexes, and atomic version updates. Run real database concurrency checks.'),('9 and 10','Security testing and delivery','Verify two-user isolation and session revocation. Run API tests and a production build, and inspect browser failure recovery.'),('11 and 12','Design company guidance and mocks','Trace your project, discuss one alternative, rehearse truthful stories, and complete coding and design mocks.')]
for label,title,body in weeks:doc.add_heading('Weeks '+label+' '+title,2);doc.add_paragraph(body)
doc.add_page_break();doc.add_heading('Interview recall and scenario practice',1)
doc.add_paragraph('Attempt questions without opening the answer. A complete response explains the concept, gives a correct example, and states a tradeoff or failure case. For scenarios, collect evidence before proposing a fix.')
for score,meaning in [('0','Cannot explain without help'),('1','Definition only'),('2','Correct answer with example'),('3','Correct answer with tradeoff and failure case')]:doc.add_paragraph(score+' '+meaning)
doc.add_heading('Review schedule',2);doc.add_paragraph('Revisit weak answers after 1, 3, 7, and 14 days. A review is useful when you retrieve and apply the idea rather than reread it. Use stable question ids in the spreadsheet.')
doc.add_heading('Scenario worksheet',2)
for label in ['Scenario id','Observed symptom','Evidence to collect','Likely cause and why','Proposed fix','Regression check','Tradeoff or remaining limitation']:doc.add_paragraph(label+' ______________________________')
doc.add_page_break();doc.add_heading('Machine coding review',1)
doc.add_paragraph('Clarify acceptance, model the state, and finish one core flow before visual polish. Spend about 10 percent of the time clarifying, 55 percent building, 20 percent checking boundaries, and 15 percent demonstrating and explaining.')
for label,points,evidence in [('Required behavior',30,'Demonstrate requested interactions and resulting state'),('Boundaries and recovery',20,'Invalid input, empty data, errors, and races where relevant'),('State and API design',15,'Clear owners, stable ids, explicit transitions, bounded work'),('Accessibility',15,'Labels, keyboard, focus, and error announcements'),('Verification',10,'Checks with independent expected outcomes'),('Explanation',10,'Invariant, complexity, tradeoff, omissions, next steps')]:doc.add_heading(label+' '+str(points)+' points',2);doc.add_paragraph(evidence)
doc.add_paragraph('Suggested self-assessment target is 75 out of 100 with no missing core flow or authorization failure. This is not a company hiring cutoff.')
doc.add_page_break();doc.add_heading('Project and company preparation',1)
doc.add_paragraph('Use official company process references and the actual role description. Confirm rounds and permitted tools with the recruiter. Repository mocks are recommendations, not predictions of company questions.')
doc.add_heading('Project defense',2)
for line in ['Describe the user problem and your actual contribution.','Trace a write from form state through validation, session, ownership, database, response, and recovery.','Demonstrate one concurrency or failure case and the test that establishes it.','Explain the chosen data model, index, and one alternative.','State the setup requirements and remaining limitations.']:doc.add_paragraph(line,style='List Bullet')
doc.add_heading('Behavioral evidence',2);doc.add_paragraph('Prepare truthful stories for ownership, disagreement, learning, failure, and ambiguity. State your responsibility and action and the actual outcome. Do not invent users, revenue, or performance measurements.')
doc.add_heading('Questions to ask',2);doc.add_paragraph('Ask how the team defines quality, how ownership is shared, what the role is expected to deliver, and how operational work is handled.')
doc.add_page_break();doc.add_heading('Weekly evidence and next actions',1)
for line in ['Week and focus','Question ids reviewed','Timed exercise and duration','Acceptance checks passed','Observed failure or unclear explanation','Project change implemented','Test or demo evidence','One next action and review date']:doc.add_paragraph(line+' ______________________________')
doc.add_heading('Useful starting references',2)
for name,url in [('MDN curriculum','https://developer.mozilla.org/en-US/docs/Learn_web_development'),('React Learn','https://react.dev/learn'),('Full Stack Open','https://fullstackopen.com/en/'),('MongoDB manual','https://www.mongodb.com/docs/manual/'),('OWASP cheat sheets','https://cheatsheetseries.owasp.org/')]:doc.add_paragraph(name+' '+url)
doc.add_paragraph('Source repository https://github.com/venkatesh7975/NEC_MERN_Notes_3rd_year')
footer=sec.footer.paragraphs[0];footer.text='MERN interview study plan | 2026-10-08 | '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(out/'mern-study-plan.docx')
print('Created complete PDF handbook and six-page editable study plan.')
